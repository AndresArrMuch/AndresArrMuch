"""Compara main y rama del Libro Diario. Uso: python jd_compare.py <main.json> <rama.json>"""
import json, sys
from openpyxl import load_workbook
m, r = json.load(open(sys.argv[1])), json.load(open(sys.argv[2]))
print('Dashboard:', m['dashboard'], '| rama:', r['dashboard'], '| igual:', m['dashboard'] == r['dashboard'])
print('Muestra formatJournalNumber de main coincide con su agrupacion:', m.get('muestra_formatJournalNumber_ok'))
dif = [k for k in m['numeros'] if m['numeros'][k] != r['numeros'].get(k)]
print('Numero visible igual en los', len(m['numeros']), 'asientos:', not dif and len(m['numeros']) == len(r['numeros']), '(distintos:', len(dif), ')', [(k, m['numeros'][k], r['numeros'].get(k)) for k in dif[:3]])
meses = sorted(set(m['por_mes']) | set(r['por_mes']))
mal = [x for x in meses if m['por_mes'].get(x) != r['por_mes'].get(x)]
print('Totales por periodo iguales en', len(meses), 'meses:', not mal, mal[:3] and [(x, m['por_mes'].get(x), r['por_mes'].get(x)) for x in mal[:3]])
print('Filtros (asientos encontrados main / rama):', [(len(a), len(b), a == b) for a, b in zip(m['filtros'], r['filtros'])])
if 'paginacion' in r: print('Paginacion:', r['paginacion']); print('Orden:', json.dumps(r['orden'], ensure_ascii=False)); print('Monto:', r['filtro_monto'])
for fm, fr in zip(m['excel'], r['excel']):
    vm = [list(x) for x in load_workbook(fm).active.iter_rows(values_only=True)]
    vr = [list(x) for x in load_workbook(fr).active.iter_rows(values_only=True)]
    print('Excel', fm.split('jd_main_')[1], 'filas', len(vm), len(vr), 'identico:', vm == vr)
print('Estados Financieros identicos:', m['estados_financieros'] == r['estados_financieros'], len(m['estados_financieros']))
for k in m['otras']: print('Pantalla', k, 'igual:', m['otras'][k] == r['otras'][k])
print('Errores JS main/rama:', m['errores_js'], r['errores_js'])
