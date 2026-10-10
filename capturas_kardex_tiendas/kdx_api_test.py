"""Kardex para encargados de tienda: lo que devuelve el servidor a cada usuario (main y rama).
Uso: python kdx_api_test.py <backend_dir> <db> <salida.json>
"""
import os, sys, json, io
BACKEND, DB, OUT = sys.argv[1:4]
os.environ['DATABASE_URL'] = f'postgresql://postgres@/{DB}?host=/var/run/postgresql&port=5433'
sys.path.insert(0, BACKEND); os.chdir(BACKEND)
import warnings; warnings.filterwarnings('ignore')
import app as _a; _a._restore_sqlite_database_from_respaldo = lambda app: None
from app import create_app
from app.models import InventoryWarehouse, InventoryItem
from app.models.user import User
from flask_jwt_extended import create_access_token
from openpyxl import load_workbook

FIKA = '976b2dc9-1a6c-4c6f-ba0e-a32b59ceb433'; FLU = 'ff8936f6-245c-4ff4-bffa-78b8fc3f8c78'
app = create_app(); cl = app.test_client()
with app.app_context():
    WH = {w.code: w.id for w in InventoryWarehouse.query.all()}
    CODE = {v: k for k, v in WH.items()}
    USERS = {}
    for key, email, cid in [('nilton', 'ventas2@fikaperu.com', FIKA), ('saul_yorko', 'ventas.tda.filomena@gmail.com', FIKA),
                            ('francisco_fika', 'ventas1@fikaperu.com', FIKA), ('francisco_flumisa', 'ventas1@fikaperu.com', FLU),
                            ('admin_fika', 'admin@t.pe', FIKA), ('admin_flumisa', 'admin@t.pe', FLU),
                            ('almacenero_fika', 'central@t.pe', FIKA), ('almacenero_flumisa', 'central@t.pe', FLU),
                            ('ventas_sin_tienda', 'ventas@t.pe', FIKA), ('tdacapon', 'capon@t.pe', FIKA), ('tdamalvinas', 'malvinas@t.pe', FIKA)]:
        u = User.query.filter_by(email=email, company_id=cid).first()
        USERS[key] = (create_access_token(identity=u.id), cid)
    ITEM_FIKA = InventoryItem.query.filter_by(company_id=FIKA).order_by(InventoryItem.item_code).first().id

def get(user, url):
    tok, _ = USERS[user]
    r = cl.get(url, headers={'Authorization': 'Bearer ' + tok})
    return r.status_code, r

def jget(user, url):
    st, r = get(user, url)
    return st, (r.get_json(silent=True) or {})

out = {}
for user, (tok, cid) in USERS.items():
    res = {}
    st, d = jget(user, f'/api/inventory/{cid}/movements?per_page=20000')
    rows = (d.get('data') or {}).get('data', []) if st == 200 else []
    res['movimientos'] = {'status': st, 'almacenes': sorted({r['warehouse_code'] for r in rows}), 'cantidad': len(rows),
                          'con_costos': any(k in r for r in rows for k in ('unit_cost', 'total_cost', 'balance_cost'))}
    res['movimientos_payload'] = rows
    st, d = jget(user, f'/api/inventory/{cid}/stock-summary')
    res['stock_resumen_almacenes'] = (st, sorted(w['code'] for w in (d.get('data') or {}).get('warehouses', [])))
    st, r = get(user, f'/api/inventory/{cid}/export/multistore')
    sheets = load_workbook(io.BytesIO(r.data)).sheetnames if st == 200 else []
    res['exportar_excel_hojas'] = (st, sheets)
    out[user] = res

# Llamadas directas con un almacen no permitido
direct = {}
direct['nilton_movimientos_capon'] = get('nilton', f'/api/inventory/{FIKA}/movements?warehouse_id={WH["TDAFIKACAPON"]}')[0]
direct['nilton_movimientos_malvinas_propio'] = get('nilton', f'/api/inventory/{FIKA}/movements?warehouse_id={WH["TDAFIKAMALVINAS"]}')[0]
direct['nilton_movimientos_almfika'] = get('nilton', f'/api/inventory/{FIKA}/movements?warehouse_id={WH["ALMFIKA"]}')[0]
direct['saul_movimientos_malvinas'] = get('saul_yorko', f'/api/inventory/{FIKA}/movements?warehouse_id={WH["TDAFIKAMALVINAS"]}')[0]
direct['nilton_items_stock_capon'] = get('nilton', f'/api/inventory/{FIKA}/items?warehouse_id={WH["TDAFIKACAPON"]}')[0]
direct['nilton_item_stock_capon'] = get('nilton', f'/api/inventory/{FIKA}/items/{ITEM_FIKA}/stock?warehouse_id={WH["TDAFIKACAPON"]}')[0]
st, d = jget('nilton', f'/api/inventory/{FIKA}/items/{ITEM_FIKA}/stock')
direct['nilton_item_stock_desglose'] = (st, sorted(CODE.get(b['warehouse_id']) for b in (d.get('data') or {}).get('by_warehouse', [])))
direct['francisco_flumisa_movimientos_tdaflucapon'] = get('francisco_flumisa', f'/api/inventory/{FLU}/movements?warehouse_id={WH["TDAFLUCAPON"]}')[0]
# Francisco con su sesion de FIKA consultando FLUMISA: se usa su cuenta de FLUMISA
st, d = jget('francisco_fika', f'/api/inventory/{FLU}/movements?per_page=20000')
direct['francisco_sesion_fika_pide_flumisa'] = (st, sorted({r['warehouse_code'] for r in (d.get('data') or {}).get('data', [])}))
# Solo lectura: crear un movimiento (cuerpo incompleto: en main da 400 sin crear nada)
tok, _ = USERS['nilton']
direct['nilton_crear_movimiento'] = cl.post(f'/api/inventory/{FIKA}/movements', headers={'Authorization': 'Bearer ' + tok}, json={}).status_code
tok, _ = USERS['almacenero_fika']
direct['almacenero_crear_movimiento_incompleto'] = cl.post(f'/api/inventory/{FIKA}/movements', headers={'Authorization': 'Bearer ' + tok}, json={}).status_code

# Otros modulos de los tres encargados (deben ser identicos en main y rama)
otros = {}
for user in ('nilton', 'saul_yorko', 'francisco_fika', 'francisco_flumisa'):
    tok, cid = USERS[user]
    for name, url in [('requerimientos', f'/api/inventory/{cid}/stock-requests'), ('transferencias', f'/api/inventory/{cid}/transfers'),
                      ('ventas', f'/api/sales/{cid}?per_page=20000'), ('cotizaciones', f'/api/sales/{cid}/workflow?per_page=20000'),
                      ('items_sin_almacen', f'/api/inventory/{cid}/items?per_page=20000'), ('almacenes', f'/api/inventory/{cid}/warehouses'),
                      ('pendientes_requerimientos', f'/api/inventory/{cid}/stock-requests/pending-count')]:
        st, d = jget(user, url)
        otros[f'{user}:{name}'] = [st, d.get('data')]
json.dump({'usuarios': out, 'directo': direct, 'otros_modulos': otros}, open(OUT, 'w'), indent=1, ensure_ascii=False, default=str)
print('listo', OUT)
